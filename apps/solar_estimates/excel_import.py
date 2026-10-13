"""
Load the equipment catalogue and assumptions from the APS estimate workbook
(APS_Solar_Estimate_TEMPLATE.xlsx) — once, to fill a fresh server
(``manage.py import_estimate_template``). After that, prices are kept in
Prices & settings; there is no upload in the dashboard.

- Catalogue tab: every inverter / battery / panel / balance-of-system row
  (cost excl VAT and specs). Replaces the whole catalogue.
- Assumptions tab: tariffs, efficiencies, margin, yield profile, VAT number
  and office phone. Rows are found by their label, not a fixed cell.
"""

from decimal import Decimal, InvalidOperation

import openpyxl
from django.db import transaction

from .models import CatalogueItem, EstimateSettings


class WorkbookError(ValueError):
    pass


BOS_KEYS = {
    'BoS|Mounting': 'mounting', 'BoS|DC kit': 'dc_kit', 'BoS|AC split': 'ac_split', 'BoS|AC full': 'ac_full',
    'BoS|Gateway T1': 'gateway_t1', 'BoS|Cabinet T2': 'cabinet_t2', 'BoS|Cabling': 'cabling',
    'BoS|Labour': 'labour', 'BoS|COC': 'coc', 'BoS|SSEG': 'sseg', 'BoS|Transport': 'transport',
}
CATEGORY = {'Inverter': 'inverter', 'Battery': 'battery', 'Panel': 'panel', 'BoS': 'bos'}

# Assumptions label (start of the text in column A) → settings field.
ASSUMPTION_ROWS = [
    ('Home User – Block 1 energy', 'block1_rate'),
    ('Home User – Block 2 energy', 'block2_rate'),
    ('Block 1 threshold', 'block1_threshold'),
    ('Service & wires charge', 'service_charge'),
    ('SSEG export credit', 'export_credit'),
    ('Battery round-trip efficiency', 'battery_efficiency'),
    ('Default share of daily use in daytime', 'daytime_share'),
    ('Battery target', 'battery_night_target'),
    ('Optimal-panel rule', 'optimal_panel_share'),
    ('Panel rating', 'panel_watts'),
    ('Sweep step', 'sweep_step'),
    ('VAT rate', 'vat_rate'),
    ('Savings prudence factor', 'prudence_factor'),
    ('Target gross margin', 'target_margin'),
]
TEXT_ROWS = [('APS VAT number', 'company_vat_number'), ('APS phone', 'company_phone')]
MONTHS = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']


def _num(value):
    if value in (None, ''):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _cost(value):
    if value in (None, ''):
        return None
    try:
        return Decimal(str(value)).quantize(Decimal('0.01'))
    except InvalidOperation:
        return None


def read_workbook(file):
    try:
        wb = openpyxl.load_workbook(file, data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001
        raise WorkbookError('That file is not a readable Excel workbook (.xlsx).') from exc
    missing = {'Catalogue', 'Assumptions'} - set(wb.sheetnames)
    if missing:
        raise WorkbookError(f'The workbook has no {" or ".join(sorted(missing))} tab — use the APS estimate template.')

    items = []
    for order, row in enumerate(wb['Catalogue'].iter_rows(min_row=4, values_only=True)):
        row = list(row) + [None] * 14
        key, tier, category, brand, model, phase = row[0], row[1], row[2], row[3], row[4], row[5]
        if not model or category not in CATEGORY:
            continue
        cat = CATEGORY[category]
        code = BOS_KEYS.get(str(key), '') if cat == 'bos' else ('panel' if cat == 'panel' else '')
        if cat == 'bos' and not code:
            continue
        items.append({
            'category': cat, 'code': code, 'tier': tier if tier in ('Tier 1', 'Tier 2') else 'Any',
            'brand': (brand or '')[:60], 'model': str(model)[:150],
            'phase': phase if phase in ('Single', 'Three') else 'Any',
            'rated_kw': _num(row[6]), 'max_pv_kwp': _num(row[7]), 'usable_kwh': _num(row[8]),
            'max_modules': int(row[9]) if _num(row[9]) else None, 'unit': (row[10] or '')[:20],
            'cost_excl_vat': _cost(row[11]), 'notes': row[13] or '', 'sort_order': order,
        })
    if not any(i['category'] == 'inverter' for i in items):
        raise WorkbookError('No inverters found on the Catalogue tab.')

    values, texts, yields = {}, {}, {}
    for row in wb['Assumptions'].iter_rows(min_row=1, max_col=3, values_only=True):
        label, value, extra = (row + (None, None, None))[:3]
        if not isinstance(label, str):
            continue
        for prefix, field in ASSUMPTION_ROWS:
            if label.startswith(prefix) and _num(value) is not None:
                values[field] = _num(value)
        for prefix, field in TEXT_ROWS:
            if label.startswith(prefix) and value is not None and str(value).strip():  # blank = keep ours
                texts[field] = str(value).strip()
        if label in MONTHS and _num(value) and _num(extra) is not None:
            yields[label] = {'month': label, 'days': int(value), 'yield': _num(extra)}
    if len(yields) == 12:
        values['yield_profile'] = [yields[m] for m in MONTHS]
    if 'sweep_step' in values:
        values['sweep_step'] = int(values['sweep_step'])
    return items, {**values, **texts}


@transaction.atomic
def import_workbook(file):
    items, settings_values = read_workbook(file)
    CatalogueItem.objects.all().delete()
    CatalogueItem.objects.bulk_create(CatalogueItem(**i) for i in items)
    s = EstimateSettings.load()
    for field, value in settings_values.items():
        setattr(s, field, value)
    s.save()
    return {'catalogue_items': len(items), 'settings_updated': sorted(settings_values)}


def read_inputs(file):
    """Proposal inputs from a filled-in estimate workbook's Input tab (e.g.
    "APS-EST 22 Chateau Circle Constantia.xlsx") — the same fields as the form."""
    wb = openpyxl.load_workbook(file, data_only=True, read_only=True)
    if 'Input' not in wb.sheetnames:
        raise WorkbookError('The workbook has no Input tab.')
    cells = {}
    for row in wb['Input'].iter_rows(min_row=1, max_row=67, max_col=4):
        for c in row:
            if c.value is not None:
                cells[c.coordinate] = c.value
    v = cells.get

    def text(ref):
        value = v(ref)
        return '' if value is None else str(value).strip()

    def iso(value):
        return value.date().isoformat() if hasattr(value, 'date') else (value.isoformat() if hasattr(value, 'isoformat') else None)

    scope = text('B33')
    readings = []
    for r in range(52, 68):
        if v(f'A{r}') is not None and _num(v(f'B{r}')) is not None:
            readings.append({'date': iso(v(f'A{r}')), 'reading': _num(v(f'B{r}')),
                             'type': 'Estimate' if text(f'C{r}') == 'Estimate' else 'Actual', 'source': text(f'D{r}')})
    overrides = {k: _num(v(ref)) for k, ref in (('daily_kwh', 'B40'), ('panels', 'B41'), ('battery_kwh', 'B42'),
                                               ('full_min_inverter_kw', 'B43'), ('split_min_inverter_kw', 'B44'))
                 if _num(v(ref)) is not None}
    if 'panels' in overrides:
        overrides['panels'] = int(overrides['panels'])
    deposit = _num(v('B47'))
    return {
        'estimate_number': text('B5') if text('B5') != '[confirm]' else '',
        'estimate_date': iso(v('B6')),
        'validity_days': int(_num(v('B7')) or 7),
        'prepared_by': text('B8'),
        'client_name': text('B12'), 'contact_person': text('B13'), 'phone': text('B14'), 'email': text('B15'),
        'client_type': text('B16'), 'site_address': text('B17'), 'erf_number': text('B18'),
        'municipal_account': text('B19'), 'tariff': text('B20') or 'Home User',
        'supply_phase': text('B23') if text('B23') in ('Single', 'Three') else 'Unknown',
        'unknown_phase_as': text('B24') if text('B24') in ('Single', 'Three') else 'Three',
        'main_breaker_a': int(_num(v('B25'))) if _num(v('B25')) else None,
        'roof_type': text('B26'), 'orientation': text('B27'), 'max_panels': int(_num(v('B28')) or 60),
        'shading': text('B29'), 'client_priority': text('B30'),
        'existing_backup': text('B31') or 'None', 'existing_backup_details': text('B32'),
        'backup_scope': 'full' if scope.startswith('Full') else 'split',
        'essential_share': _num(v('B34')) or 0.5,
        'export_to_grid': text('B35') != 'N',
        'daytime_share_override': _num(v('B36')),
        'site_notes': text('B37'),
        'overrides': overrides,
        'deposit_percent': deposit * 100 if deposit is not None and deposit <= 1 else deposit,
        'balance_terms': text('B48'),
        'readings': readings,
    }
